# -*- coding: utf-8 -*-
"""
納瓦爾工作坊 — 完整版 PPTX 生成器
從 slides-master.md 解析（單一來源），輸出那瓦爾藍金設計的完整投影片。
跑法： uv run --with python-pptx python build_pptx.py
"""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette（那瓦爾藍金）----------
NAVY=RGBColor(0x08,0x27,0x43); CARD=RGBColor(0x0c,0x32,0x58); INSET=RGBColor(0x06,0x24,0x3f)
GOLD=RGBColor(0xE9,0xB6,0x27); GOLDS=RGBColor(0xf4,0xd7,0x6a); GOLDD=RGBColor(0xb9,0x8a,0x16)
WHITE=RGBColor(0xff,0xff,0xff); INK=RGBColor(0xf1,0xea,0xd9); MUTED=RGBColor(0xae,0xbc,0xcb)
FAINT=RGBColor(0x7e,0x93,0xa8); LINE=RGBColor(0x1f,0x4a,0x78)
SERIF="Noto Serif TC"; SANS="Noto Sans TC"; PLAY="Playfair Display"

W=Inches(13.333); H=Inches(7.5)
prs=Presentation(); prs.slide_width=W; prs.slide_height=H
BLANK=prs.slide_layouts[6]

TABS={
 'home':'首頁','p1':'PART I · 39 金句','p2':'PART II · 孟修 WIKI 大腦',
 'p3':'PART III · 為什麼工作坊','p4':'PART IV · 造運引擎','p5':'PART V · 將自己商品化',
 'p6':'PART VI · 顧問團','c21':'🔥 21 天挑戰','tl':'⏳ 時間槓桿','lib':'📚 圖書館','man':'📋 使用手冊',
}

def setfont(run,name,size,color,bold=False,italic=False):
    f=run.font; f.size=Pt(size); f.bold=bold; f.italic=italic; f.color.rgb=color; f.name=name
    rPr=run._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set('typeface',name)

def slide():
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,W,H)
    r.fill.solid(); r.fill.fore_color.rgb=NAVY; r.line.fill.background(); r.shadow.inherit=False
    return s

def rect(s,l,t,w,h,fill=None,line=None,lw=1.0,rounded=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l),Inches(t),Inches(w),Inches(h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    shp.shadow.inherit=False; return shp

def text(s,l,t,w,h,paras,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,pa in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=pa.get('align',PP_ALIGN.LEFT)
        p.space_after=Pt(pa.get('after',4)); p.space_before=Pt(pa.get('before',0))
        if pa.get('line'): p.line_spacing=pa['line']
        for (txt,font,size,color,bold,italic) in pa['r']:
            run=p.add_run(); run.text=txt; setfont(run,font,size,color,bold,italic)
    return tb

# ---------- parse master ----------
def parse(path):
    slides=[]; cur=None
    for raw in open(path,encoding='utf-8'):
        ln=raw.rstrip('\n')
        m=re.match(r'^## \[([a-z0-9]+)-(\d+)\]\s*(.+)$',ln)
        if m:
            if cur: slides.append(cur)
            cur={'tab':m.group(1),'idx':int(m.group(2)),'title':m.group(3).strip(),'point':'','body':[]}
        elif cur is not None:
            mp=re.match(r'^\*\*重點\*\*：(.+)$',ln)
            if mp: cur['point']=mp.group(1).strip()
            elif ln.startswith('- '): cur['body'].append(ln[2:].strip())
    if cur: slides.append(cur)
    return slides

# ---------- slide renderers ----------
def is_cover(t): return '封面' in t
def is_divider(t): return t.startswith('分類') or '總覽' in t

def render_cover(sld):
    s=slide()
    text(s,1,1.0,11.3,0.5,[{'r':[(TABS.get(sld['tab'],''),PLAY,15,GOLD,True,False)],'align':PP_ALIGN.CENTER}])
    title=sld['title'].replace('封面：','').replace('封面','').strip('： ')
    text(s,1,2.3,11.3,1.6,[{'r':[(title or TABS.get(sld['tab'],''),SERIF,46,GOLDS,True,False)],'align':PP_ALIGN.CENTER}])
    if sld['point']:
        rect(s,1.6,4.2,10.1,1.0,fill=CARD,line=GOLDD,lw=1.2)
        text(s,1.9,4.45,9.5,0.6,[{'r':[(sld['point'],SERIF,19,GOLD,True,False)],'align':PP_ALIGN.CENTER}])
    body=sld['body']
    if body:
        text(s,1.4,5.45,10.5,1.4,[{'r':[(b,SANS,13.5,MUTED,False,False)],'align':PP_ALIGN.CENTER,'line':1.3} for b in body[:4]])
    return s

def render_content(sld,n):
    s=slide()
    rect(s,0.6,0.62,0.55,0.07,fill=GOLD,rounded=False)
    text(s,0.6,0.42,11.6,0.35,[{'r':[(TABS.get(sld['tab'],''),PLAY,12.5,GOLD,True,False)]}])
    # title
    text(s,0.6,0.82,12.1,0.95,[{'r':[(sld['title'],SERIF,25,WHITE,True,False)]}])
    top=1.95
    # 重點 highlight box
    if sld['point']:
        rect(s,0.6,top,12.13,0.78,fill=INSET,line=GOLDD,lw=1.1)
        text(s,0.85,top+0.14,11.6,0.5,[{'r':[('重點　',PLAY,12,GOLDD,True,False),(sld['point'],SERIF,15.5,GOLDS,True,False)],'line':1.15}],anchor=MSO_ANCHOR.MIDDLE)
        top+=1.0
    # body bullets — auto size
    body=sld['body']
    bsz=14.5 if len(body)<=4 else (12.5 if len(body)<=6 else 11)
    paras=[]
    for b in body:
        paras.append({'r':[('▸ ',SANS,bsz,GOLD,True,False),(b,SANS,bsz,INK,False,False)],'line':1.25,'after':6})
    if paras: text(s,0.7,top+0.05,12.0,6.4-top,paras)
    # footer
    text(s,0.6,7.0,9,0.3,[{'r':[('納瓦爾工作坊 · 將個人商品化 · 給扛霸子的致富之道',SANS,9,FAINT,False,False)]}])
    text(s,12.2,7.0,0.9,0.3,[{'r':[(str(n).zfill(3),PLAY,11,FAINT,False,False)],'align':PP_ALIGN.RIGHT}])
    return s

# ---------- build ----------
HERE=os.path.dirname(os.path.abspath(__file__))
slides=parse(os.path.join(HERE,'slides-master.md'))
n=0
for sld in slides:
    n+=1
    if is_cover(sld['title']): render_cover(sld)
    else: render_content(sld,n)
out=os.path.join(HERE,'納瓦爾工作坊-完整版-投影片.pptx')
prs.save(out)
print("OK ->",out,"| slides:",len(prs.slides._sldIdLst))
